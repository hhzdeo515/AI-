"""Offline graph, bounded repair, retrieval provenance and HTTP contract tests."""
import io
import json
import sys
import tempfile
import threading
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from lg_assistant import config, exam_knowledge as kb, graph, store, vision
from lg_assistant.exam_graph import EXAM_GRAPH, options
from lg_assistant.web.app import create_app


class ExamTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.saved = {k: getattr(config, k) for k in ("DATA_DIR", "DB_PATH", "CHECKPOINT_DB", "EXPORT_DIR", "UPLOAD_DIR", "ACCESS_TOKEN")}
        root = Path(self.temp.name)
        config.DATA_DIR = root
        config.DB_PATH = root / "app.sqlite3"
        config.CHECKPOINT_DB = root / "checkpoints.sqlite3"
        config.EXPORT_DIR = root / "exports"
        config.UPLOAD_DIR = root / "uploads"
        config.ACCESS_TOKEN = ""
        store._local = threading.local()
        store._initialised = False
        self.image = root / "question.png"
        self.image.write_bytes(b"mock image")
        self.draft = {"answerable": True, "candidate": "B", "calculations": [{"expression": "8/2"}]}
        self.final = {"answerable": True, "answer": "B，4", "explanation": "8/2=4", "repair_issues": []}
        # Main-graph/HTTP tests now inventory the photo before the existing
        # single-question mocks run. Pixel cropping has dedicated batch tests.
        self.batch_stubs = ExitStack()
        self.batch_stubs.enter_context(patch("lg_assistant.exam_batch.detect_questions", return_value={
            "questions": [{"number": "1", "preview": "", "complete": True, "missing": [],
                           "regions": [{"page": 1, "bbox": [0, 0, 1000, 1000]}], "context_regions": []}],
            "notes": []}))
        self.batch_stubs.enter_context(patch("lg_assistant.exam_batch._crop_images",
            side_effect=lambda images, question, crop_dir: list(images)))
        self.batch_stubs.enter_context(patch("lg_assistant.llm.json_chat", side_effect=AssertionError("unexpected classification network call")))
        self.batch_stubs.enter_context(patch("lg_assistant.jev.evaluate", side_effect=AssertionError("unexpected JEV network call")))

    def tearDown(self):
        self.batch_stubs.close()
        conn = getattr(store._local, "conn", None)
        if conn: conn.close()
        store._local = threading.local()
        store._initialised = False
        for k, v in self.saved.items(): setattr(config, k, v)
        self.temp.cleanup()

    def run_exam(self, *, observation="题干：计算8/2。A 2 B 4 C 6 D 8", event=None, reviewer=None, solver=None):
        with patch("lg_assistant.vision.observe", return_value=observation), \
             patch("lg_assistant.vision.solve", side_effect=solver or (lambda *a, **k: dict(self.draft))) as solve, \
             patch("lg_assistant.vision.review", side_effect=reviewer or (lambda *a, **k: dict(self.final))) as review, \
             patch("lg_assistant.search.search_answer", side_effect=AssertionError("unexpected web call")):
            state = EXAM_GRAPH.invoke({"owner": "alice", "files": [str(self.image)], "text": "解题", "event": {"exam": event or {}}})
            return state, solve.call_count, review.call_count

    def test_standard_route_one_specialist_no_loop(self):
        out, solves, reviews = self.run_exam()
        self.assertEqual((solves, reviews), (1, 1))
        self.assertEqual(out["specialist"], "quantitative")
        self.assertEqual(out["tools"]["calculations"][0]["result"], 4)
        self.assertEqual(out["repair_count"], 0)
        self.assertTrue(all(t["ms"] >= 0 for t in out["trace"]))

    def test_specific_repair_rechecks_and_stops(self):
        finals = iter([{"answerable": False, "repair_issues": [{"code": "option_mapping", "detail": "B与C标签错位，重新看原图"}]}, self.final])
        out, solves, reviews = self.run_exam(reviewer=lambda *a, **k: next(finals))
        self.assertEqual((solves, reviews), (2, 2))
        self.assertTrue(out["final"]["answerable"])
        self.assertEqual(out["repair_count"], 1)

    def test_repeated_issue_hits_cap_and_withholds_answer(self):
        self.final.update(repair_issues=[{"code": "calculation", "detail": "关键分母仍有冲突"}])
        out, solves, _ = self.run_exam()
        self.assertEqual(solves, 2)
        self.assertFalse(out["final"]["answerable"])
        self.assertEqual(out["final"]["answer"], "")
        self.assertNotIn("答案是 B", out["result"]["speech"])

    def test_missing_image_detail_does_not_loop(self):
        self.final = {"answerable": False, "needed": "请补拍右侧选项", "speech": "答案是B", "repair_issues": []}
        out, solves, _ = self.run_exam()
        self.assertEqual(solves, 1)
        self.assertNotIn("答案是B", out["result"]["speech"])

    def test_tool_error_cannot_be_overruled_by_model(self):
        self.draft["calculations"] = [{"expression": "1/0"}]
        out, solves, _ = self.run_exam()
        self.assertEqual(solves, 2)
        self.assertFalse(out["final"]["answerable"])

    def test_internal_missing_library_stops_before_solver(self):
        out, solves, _ = self.run_exam(event={"profile": "internal"})
        self.assertEqual(solves, 0)
        self.assertIn("对应制度", out["result"]["text"])

    def add_reference(self):
        return kb.add("alice", "培训考核办法", "员工培训考核成绩达到八十分为合格。补考安排由培训部门另行通知。", "2025版")["id"]

    def test_chinese_retrieval_owner_scope_and_version(self):
        doc = self.add_reference()
        rows = kb.retrieve("alice", "员工培训考核成绩多少分合格", [doc])
        self.assertEqual(rows[0]["version"], "2025版")
        with self.assertRaises(ValueError): kb.retrieve("bob", "培训考核", [doc])
        self.assertEqual(kb.documents("bob"), [])

    def test_identical_import_is_idempotent(self):
        self.assertEqual(self.add_reference(), self.add_reference())
        self.assertEqual(len(kb.documents("alice")), 1)

    def test_source_quote_is_verified_and_rendered(self):
        doc = self.add_reference()
        self.final.update(answer="八十分", citations=[{"id": doc + ":1", "quote": "员工培训考核成绩达到八十分为合格。"}])
        out, solves, _ = self.run_exam(observation="员工培训考核成绩达到多少分合格？", event={"profile": "internal", "document_ids": [doc]})
        self.assertEqual(solves, 1)
        self.assertIn("2025版", out["result"]["text"])
        self.assertTrue(out["final"]["answerable"])

    def test_fabricated_quote_is_not_an_answer(self):
        doc = self.add_reference()
        self.final["citations"] = [{"id": doc + ":1", "quote": "六十分合格"}]
        out, solves, _ = self.run_exam(observation="员工培训考核成绩达到多少分合格？", event={"document_ids": [doc]})
        self.assertFalse(out["final"]["answerable"])
        self.assertEqual(solves, 1)

    def test_unselected_sources_are_never_used(self):
        self.add_reference()
        self.assertEqual(kb.retrieve("alice", "培训考核", []), [])

    def test_delete_removes_chunks_and_fts(self):
        doc = self.add_reference()
        self.assertFalse(kb.remove("bob", doc))
        self.assertTrue(kb.remove("alice", doc))
        with kb.connection() as db:
            self.assertEqual(db.execute("SELECT count(*) FROM exam_chunks").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT count(*) FROM exam_fts").fetchone()[0], 0)

    def test_stale_selection_cannot_fall_back_to_memory(self):
        out, solves, _ = self.run_exam(event={"document_ids": ["f" * 32]})
        self.assertEqual(solves, 0)
        self.assertIn("重新选择", out["result"]["text"])

    def test_unrelated_material_stops_early(self):
        doc = self.add_reference()
        out, solves, _ = self.run_exam(observation="三角函数周期求解", event={"document_ids": [doc]})
        self.assertEqual(solves, 0)
        self.assertIn("没有检索到", out["result"]["text"])

    def test_internal_and_selected_material_disable_web(self):
        self.assertFalse(options({"exam": {"profile": "internal", "allow_web": True}})["allow_web"])
        self.assertFalse(options({"exam": {"document_ids": ["f" * 32], "allow_web": True}})["allow_web"])

    def test_time_sensitive_with_web_disabled_stops(self):
        out, solves, _ = self.run_exam(observation="党的二十届三中全会是否召开？", event={"allow_web": False})
        self.assertEqual(solves, 0)
        self.assertIn("核实依据", out["result"]["text"])

    def test_malformed_answer_flag_is_rejected(self):
        self.final["answerable"] = "false"
        out, _, _ = self.run_exam()
        self.assertFalse(out["result"]["artifacts"][0]["answerable"])

    def test_graph_registered_as_nested_subgraph(self):
        main = graph.build_graph(with_telemetry=False)
        self.assertIn("exam_vision", dict(main.get_subgraphs()))
        self.assertIn("review_agent", EXAM_GRAPH.get_graph().nodes)

    def test_public_catalog_has_scope_without_private_documents(self):
        self.add_reference()
        client=create_app(app=object()).test_client()
        catalog=client.get('/api/exam/public-knowledge').json
        self.assertEqual(len(catalog['concepts']),catalog['count'])
        self.assertTrue(all(note['content'] and note['version_label'] for note in catalog['concepts']))
        self.assertEqual(catalog['count'],248)
        self.assertEqual(catalog['modules']['图形推理'],30)
        self.assertNotIn('员工培训考核成绩',json.dumps(catalog,ensure_ascii=False))

    def test_http_roundtrip_and_owner_isolation(self):
        client = create_app(app=graph.build_graph(with_telemetry=False)).test_client()
        r = client.post("/api/exam/knowledge", json={"owner": "alice", "title": "制度", "content": "员工培训考核成绩达到八十分为合格。", "version": "2025"})
        self.assertEqual(r.status_code, 200)
        doc = r.json["id"]
        self.assertEqual(client.get(f"/api/exam/knowledge/{doc}?owner=bob").status_code, 404)
        self.assertEqual(client.get(f"/api/exam/knowledge/{doc}?owner=alice").json["version"], "2025")
        self.assertEqual(client.delete(f"/api/exam/knowledge/{doc}?owner=bob").status_code, 404)
        self.assertEqual(client.delete(f"/api/exam/knowledge/{doc}?owner=alice").status_code, 200)

    def test_http_file_import_and_invalid_format(self):
        client = create_app(app=graph.build_graph(with_telemetry=False)).test_client()
        r = client.post("/api/exam/knowledge", data={"title": "课程", "file": (io.BytesIO("数学课程讲义".encode()), "course.txt")})
        self.assertEqual(r.status_code, 200)
        r = client.post("/api/exam/knowledge", data={"title": "课程", "file": (io.BytesIO(b"x"), "course.pdf")})
        self.assertEqual(r.status_code, 400)

    def test_knowledge_api_uses_existing_auth_guard(self):
        config.ACCESS_TOKEN = "test-secret"
        client = create_app(app=graph.build_graph(with_telemetry=False)).test_client()
        self.assertEqual(client.get("/api/exam/knowledge").status_code, 401)

    def test_web_evidence_requires_sources(self):
        with patch("lg_assistant.vision.observe", return_value="党的二十届三中全会是否召开？"), \
             patch("lg_assistant.search.search_answer", return_value={"answer": "mock fact", "sources": []}), \
             patch("lg_assistant.vision.solve") as solve:
            out = EXAM_GRAPH.invoke({"files": [str(self.image)], "event": {}})
        self.assertFalse(solve.called)
        self.assertIn("可追溯来源", out["result"]["text"])

    def test_main_checkpoint_resets_private_exam_state(self):
        from langgraph.checkpoint.memory import InMemorySaver
        main = graph.build_graph(checkpointer=InMemorySaver(), with_telemetry=False)
        cfg = graph.run_config("alice", "exam-test")
        doc = self.add_reference()
        with patch("lg_assistant.vision.observe", return_value="员工培训考核成绩多少分为合格"), \
             patch("lg_assistant.vision.solve", return_value=self.draft), \
             patch("lg_assistant.vision.review", return_value={**self.final, "citations": [{"id": doc + ":1", "quote": "员工培训考核成绩达到八十分为合格。"}]}):
            first = main.invoke({"owner": "alice", "text": "解题", "scene_hint": "exam", "files": [str(self.image)], "event": {"exam": {"document_ids": [doc]}, "practice": {"agent": "ability"}}}, cfg)
        self.assertTrue(first["result"]["artifacts"][0]["reference_ids"])
        with patch("lg_assistant.vision.observe", return_value="计算8/2"), \
             patch("lg_assistant.vision.solve", return_value=self.draft), \
             patch("lg_assistant.vision.review", return_value=dict(self.final)):
            second = main.invoke({"owner": "alice", "text": "重新解题", "scene_hint": "exam", "files": [str(self.image)], "event": {"practice": {"agent": "ability"}}}, cfg)
        self.assertEqual(second["result"]["artifacts"][0]["reference_ids"], [])
        third = main.invoke({"owner": "alice", "text": "解题", "scene_hint": "exam", "files": [str(self.image)], "event": {"exam": {"profile": []}, "practice": {"agent": "ability"}}}, cfg)
        failed_batch = next(a for a in third["result"]["artifacts"] if a["kind"] == "exam_batch")
        self.assertEqual(failed_batch["answered_count"], 0)
        self.assertEqual(failed_batch["questions"], [])
        self.assertNotIn("培训考核办法", third["result"]["text"])

    def test_chat_form_uses_selected_knowledge_and_archives_citation(self):
        doc = self.add_reference()
        final = {**self.final, "citations": [{"id": doc + ":1", "quote": "员工培训考核成绩达到八十分为合格。"}]}
        client = create_app(app=graph.build_graph(with_telemetry=False)).test_client()
        with patch("lg_assistant.vision.observe", return_value="员工培训考核成绩多少分为合格"), \
             patch("lg_assistant.vision.solve", return_value=self.draft), \
             patch("lg_assistant.vision.review", return_value=final), \
             patch("lg_assistant.search.search_answer", side_effect=AssertionError("private source leaked to search")):
            r = client.post("/api/chat", data={"owner": "alice", "scene": "exam", "text": "解题", "event": json.dumps({"exam": {"profile": "internal", "document_ids": [doc]}, "practice": {"agent": "ability"}}), "files": (io.BytesIO(b"mock"), "question.png")})
        self.assertEqual(r.status_code, 200)
        self.assertIn("2025版", r.json["text"])
        self.assertTrue(r.json["archived_id"])
        self.assertTrue(r.json["artifacts"][0]["reference_ids"])

    def test_word_import_and_corrupt_document(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("单位内部培训管理制度")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "合格标准"
        table.cell(0, 1).text = "八十分"
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        client = create_app(app=graph.build_graph(with_telemetry=False)).test_client()
        r = client.post("/api/exam/knowledge", data={"title": "制度", "file": (buf, "rules.docx")})
        self.assertEqual(r.status_code, 200)
        content = client.get("/api/exam/knowledge/" + r.json["id"]).json["content"]
        self.assertIn("合格标准 | 八十分", content)
        r = client.post("/api/exam/knowledge", data={"title": "损坏", "file": (io.BytesIO(b"broken"), "bad.docx")})
        self.assertEqual(r.status_code, 400)


if __name__ == "__main__":
    unittest.main(verbosity=2)
