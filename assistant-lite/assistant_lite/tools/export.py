"""资料导出：Markdown / 纯文本 / JSON / CSV / Word / PDF。

改写自旧 dify-assistant/service.py 的 export_resources（其中文排版处理已在本机验证）。
差异：PDF 用 reportlab 内置的 STSong-Light 中文字体，不依赖外部字体文件，
避免把 9.7MB 的 simhei.ttf 塞进仓库。
"""

from __future__ import annotations

import csv
import json
import re
import uuid
from pathlib import Path
from typing import Any, Iterator

from .. import config

FORMATS = ("md", "txt", "json", "csv", "docx", "pdf")

CJK_FONT = "Microsoft YaHei"
PDF_CJK_FONT = "STSong-Light"


class ExportError(ValueError):
    """导出参数或内容有问题。"""


# --------------------------------------------------------------------------- #
# Markdown 解析
# --------------------------------------------------------------------------- #
def clean_text(text: str) -> str:
    """去掉 Markdown 标记与 LaTeX 残留，让纯文本/Word/PDF 里可读。"""
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text or "")
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("\\times", " × ").replace("\\div", " ÷ ")
    for marker in ("\\[", "\\]", "\\(", "\\)"):
        text = text.replace(marker, "")
    return text


def blocks(text: str) -> Iterator[tuple[str, Any]]:
    """把 Markdown 拆成 (heading|table|paragraph, 内容) 序列。"""
    lines = (text or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line or line in ("---", "-"):
            continue
        if (
            line.startswith("|")
            and i < len(lines)
            and re.match(r"^\|?[\s:|\-]+\|?$", lines[i].strip())
        ):
            rows = [[clean_text(x.strip()) for x in line.strip("|").split("|")]]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(
                    [clean_text(x.strip()) for x in lines[i].strip().strip("|").split("|")]
                )
                i += 1
            yield "table", rows
        elif line.startswith("#"):
            yield "heading", clean_text(line.lstrip("#").strip())
        else:
            yield "paragraph", clean_text(line)


# --------------------------------------------------------------------------- #
# 文本类
# --------------------------------------------------------------------------- #
def render_markdown(rows: list[dict[str, Any]]) -> str:
    return "\n\n---\n\n".join(f"# {r.get('title', '资料')}\n\n{r.get('content', '')}" for r in rows)


def render_text(rows: list[dict[str, Any]]) -> str:
    out = []
    for r in rows:
        out.append(r.get("title", "资料"))
        out.append("=" * 40)
        for kind, value in blocks(r.get("content", "")):
            if kind == "heading":
                out.append(f"\n【{value}】")
            elif kind == "table":
                out.extend("  ".join(row) for row in value)
            else:
                out.append(value)
        out.append("")
    return "\n".join(out)


def _write_json(path: Path, rows: list[dict[str, Any]]) -> None:
    payload = [
        {
            k: r.get(k)
            for k in ("id", "scene", "title", "content", "source", "created")
        }
        for r in rows
    ]
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    # utf-8-sig 让 Excel 正确识别中文
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["资料ID", "类型", "标题", "创建时间", "完整内容"])
        for r in rows:
            w.writerow(
                [
                    r.get("id", ""),
                    r.get("scene", ""),
                    r.get("title", ""),
                    r.get("created", ""),
                    r.get("content", ""),
                ]
            )


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #
def _write_docx(path: Path, rows: list[dict[str, Any]]) -> None:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
    except ImportError as e:  # pragma: no cover
        raise ExportError(f"缺少 python-docx：{e}") from e

    d = Document()

    def set_cjk(style) -> None:
        style.font.name = CJK_FONT
        style.font.size = Pt(10.5)
        rpr = style.element.get_or_add_rPr()
        rpr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    set_cjk(d.styles["Normal"])
    for name in ("Title", "Heading 1", "Heading 2"):
        st = d.styles[name]
        set_cjk(st)
        st.font.color.rgb = RGBColor(0, 0, 0)

    for r in rows:
        d.add_heading(r.get("title", "资料"), 0)
        for kind, value in blocks(r.get("content", "")):
            if kind == "heading":
                d.add_heading(value, 2)
            elif kind == "table":
                if not value:
                    continue
                width = max(len(x) for x in value)
                table = d.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for row in value:
                    cells = table.add_row().cells
                    for j, cell in enumerate(row):
                        if j < width:
                            cells[j].text = cell
            else:
                d.add_paragraph(value)
        d.add_page_break()
    d.save(str(path))


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _xml_escape(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _write_pdf(path: Path, rows: list[dict[str, Any]]) -> None:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import (
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as e:  # pragma: no cover
        raise ExportError(f"缺少 reportlab：{e}") from e

    # 内置中文 CID 字体，无需外部字体文件
    pdfmetrics.registerFont(UnicodeCIDFont(PDF_CJK_FONT))

    title_style = ParagraphStyle(
        "cn-title", fontName=PDF_CJK_FONT, fontSize=18, leading=26, spaceAfter=10
    )
    head_style = ParagraphStyle(
        "cn-head", fontName=PDF_CJK_FONT, fontSize=13, leading=20, spaceBefore=10, spaceAfter=4
    )
    body_style = ParagraphStyle(
        "cn-body", fontName=PDF_CJK_FONT, fontSize=10.5, leading=17, spaceAfter=3
    )

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=rows[0].get("title", "资料") if rows else "资料",
    )

    story: list[Any] = []
    for idx, r in enumerate(rows):
        if idx:
            story.append(PageBreak())
        story.append(Paragraph(_xml_escape(r.get("title", "资料")), title_style))
        for kind, value in blocks(r.get("content", "")):
            if kind == "heading":
                story.append(Paragraph(_xml_escape(value), head_style))
            elif kind == "table":
                if not value:
                    continue
                width = max(len(x) for x in value)
                data = [
                    [Paragraph(_xml_escape(c), body_style) for c in row]
                    + [""] * (width - len(row))
                    for row in value
                ]
                t = Table(data, hAlign="LEFT")
                t.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.4, (0.6, 0.6, 0.6)),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 4),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ]
                    )
                )
                story.append(t)
                story.append(Spacer(1, 6))
            else:
                story.append(Paragraph(_xml_escape(value), body_style))
    if not story:
        story.append(Paragraph("（空）", body_style))

    doc.build(story)


# --------------------------------------------------------------------------- #
# 对外入口
# --------------------------------------------------------------------------- #
def export_rows(
    rows: list[dict[str, Any]], fmt: str, out_dir: Path | None = None
) -> Path:
    """把若干资料导出成指定格式，返回生成的文件路径。"""
    fmt = (fmt or "").lower().strip().lstrip(".")
    if fmt not in FORMATS:
        raise ExportError(f"不支持的格式：{fmt or '(空)'}；支持 {'/'.join(FORMATS)}")
    if not rows:
        raise ExportError("没有可导出的资料。先生成会议纪要、题解或训练总结。")

    out_dir = Path(out_dir) if out_dir else config.EXPORT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{uuid.uuid4().hex}.{fmt}"

    if fmt == "md":
        path.write_text(render_markdown(rows), encoding="utf-8")
    elif fmt == "txt":
        path.write_text(render_text(rows), encoding="utf-8")
    elif fmt == "json":
        _write_json(path, rows)
    elif fmt == "csv":
        _write_csv(path, rows)
    elif fmt == "docx":
        _write_docx(path, rows)
    else:
        _write_pdf(path, rows)
    return path


def export_resource(owner: str, resource_id: str, fmt: str) -> Path:
    """按资料 ID 导出单个资料。"""
    from .. import session

    row = session.get_resource(owner, resource_id)
    if not row:
        raise ExportError(f"找不到资料：{resource_id}")
    return export_rows([row], fmt)
