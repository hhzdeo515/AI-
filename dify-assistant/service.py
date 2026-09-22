"""Local single-owner asset service. Dify calls require a dedicated bearer token."""
import os,json,sqlite3,uuid,time,datetime,hashlib,hmac,io,zipfile,ast,operator,re,html,threading,math
from pathlib import Path
from urllib.parse import urlparse
import requests
from flask import Flask,request,jsonify,send_file,abort
from exam_tools import check_grids

ROOT=Path('/data');ROOT.mkdir(exist_ok=True)
SET=json.loads(Path('/config/settings.json').read_text())
app=Flask(__name__);app.config['MAX_CONTENT_LENGTH']=24*1024*1024
BASE='http://localhost:8792';DB=ROOT/'library.sqlite3'
def conn():
    c=sqlite3.connect(DB,timeout=30);c.row_factory=sqlite3.Row;return c
with conn() as c:
    c.execute('CREATE TABLE IF NOT EXISTS resources(id TEXT PRIMARY KEY, owner TEXT, scene TEXT, title TEXT, content TEXT, source TEXT, files TEXT, created TEXT, document_id TEXT, index_status TEXT)')
    c.execute('CREATE TABLE IF NOT EXISTS sessions(owner TEXT, id TEXT, state TEXT, PRIMARY KEY(owner,id))')
    c.execute('CREATE TABLE IF NOT EXISTS receipts(owner TEXT, request_id TEXT, stage TEXT, response TEXT, PRIMARY KEY(owner,request_id,stage))')
    c.execute('CREATE TABLE IF NOT EXISTS exports(id TEXT PRIMARY KEY, owner TEXT, path TEXT, expires INTEGER)')
LOCK=threading.RLock()

def now():return datetime.datetime.now(datetime.timezone.utc).isoformat()
def dumps(x):return json.dumps(x,ensure_ascii=False)
def uid():return str(uuid.uuid4())
def load_state(owner,sid):
    with conn() as c:r=c.execute('SELECT state FROM sessions WHERE owner=? AND id=?',(owner,sid)).fetchone()
    return json.loads(r[0]) if r else {'active_scene':'general','meeting':{'status':'idle','transcript':''},'recent_resources':[],'last_question':'','last_answer':''}
def save_state(owner,sid,state):
    with conn() as c:c.execute('INSERT OR REPLACE INTO sessions VALUES(?,?,?)',(owner,sid,dumps(state)))
def receipt(owner,rid,stage):
    with conn() as c:r=c.execute('SELECT response FROM receipts WHERE owner=? AND request_id=? AND stage=?',(owner,rid,stage)).fetchone()
    return json.loads(r[0]) if r else None
def remember(owner,rid,stage,result):
    with conn() as c:c.execute('INSERT OR REPLACE INTO receipts VALUES(?,?,?,?)',(owner,rid,stage,dumps(result)))
    return result
@app.before_request
def auth():
    if request.path.startswith('/api/') and not hmac.compare_digest(request.headers.get('Authorization',''),'Bearer '+SET['service_token']):abort(401)
def payload():
    p=request.get_json(force=True)
    if not isinstance(p,dict) or not p.get('owner') or not p.get('session_id'):abort(400,'owner and session_id required')
    return p

def sync_knowledge(rid):
    try:
        with conn() as c:r=c.execute('SELECT * FROM resources WHERE id=?',(rid,)).fetchone()
        text=f"资料ID：{rid}\n时间：{r['created']}\n类型：{r['scene']}\n\n{r['content']}\n\n来源内容：\n{r['source']}"
        res=requests.post(f"http://api:5001/v1/datasets/{SET['dataset_id']}/document/create-by-text",headers={'Authorization':'Bearer '+SET['dataset_token']},json={'name':r['title']+' '+rid[:8],'text':text,'indexing_technique':'economy','process_rule':{'mode':'automatic'}},timeout=45)
        res.raise_for_status();j=res.json()
        with conn() as c:c.execute('UPDATE resources SET document_id=?,index_status=? WHERE id=?',(j['document']['id'],'submitted',rid))
    except Exception:
        with conn() as c:c.execute('UPDATE resources SET index_status=? WHERE id=?',('failed',rid))

def preserve_files(files):
    saved=[]
    for f in files[:5]:
        url=f.get('url') or f.get('remote_url') or ''
        p=urlparse(url)
        if p.hostname not in (None,'localhost','127.0.0.1','api','host.docker.internal') or not p.path.startswith('/files/'):
            saved.append({'name':f.get('filename','图片'),'status':'not_archived','reason':'需要 Dify 本机文件引用'});continue
        target='http://api:5001'+p.path+('?' + p.query if p.query else '')
        try:
            res=requests.get(target,timeout=30);res.raise_for_status()
            if len(res.content)>20*1024*1024:raise ValueError('file too large')
            ext=Path(f.get('filename','image.png')).suffix.lower()
            if ext not in ('.png','.jpg','.jpeg','.webp','.gif','.txt','.pdf'):ext='.bin'
            dest=ROOT/'originals'/ (uid()+ext);dest.parent.mkdir(exist_ok=True);dest.write_bytes(res.content)
            saved.append({'name':Path(f.get('filename','原图'+ext)).name,'path':str(dest),'status':'saved'})
        except Exception:saved.append({'name':f.get('filename','图片'),'status':'failed'})
    return saved

def archive(owner,scene,title,content,source,files):
    rid=uid();stored=preserve_files(files)
    with conn() as c:c.execute('INSERT INTO resources VALUES(?,?,?,?,?,?,?,?,?,?)',(rid,owner,scene,title[:100],content,source,dumps(stored),now(),'','pending'))
    threading.Thread(target=sync_knowledge,args=(rid,),daemon=True).start()
    return rid,stored

def resource_rows(owner,query='',ids=None):
    with conn() as c:
        if ids:
            rows=c.execute('SELECT * FROM resources WHERE owner=? AND id IN ('+','.join('?' for _ in ids)+') ORDER BY created DESC',[owner]+ids).fetchall()
            if len(rows)!=len(set(ids)):raise ValueError('资料不存在或无访问权限')
        else:
            rows=c.execute('SELECT * FROM resources WHERE owner=? ORDER BY created DESC LIMIT 100',(owner,)).fetchall()
    if query:
        terms=[x for x in re.split(r'\s+',query) if x]
        rows=[r for r in rows if all(t.lower() in (r['title']+r['content']+r['source']).lower() for t in terms)]
    return rows[:30]

def sign(eid,expiry):return hmac.new(SET['download_secret'].encode(),f'{eid}:{expiry}'.encode(),hashlib.sha256).hexdigest()
def clean_text(text):
    text=re.sub(r'\*\*(.*?)\*\*',r'\1',text)
    text=re.sub(r'`([^`]+)`',r'\1',text)
    text=text.replace('\\times',' × ').replace('\\div',' ÷ ')
    for marker in ('\\[','\\]','\\(','\\)'):text=text.replace(marker,'')
    return text
def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line or line in ('---','-'):continue
        if line.startswith('|') and i<len(lines) and re.match(r'^\|?[\s:|\-]+\|?$',lines[i].strip()):
            rows=[[clean_text(x.strip()) for x in line.strip('|').split('|')]];i+=1
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([clean_text(x.strip()) for x in lines[i].strip().strip('|').split('|')]);i+=1
            yield 'table',rows
        elif line.startswith('#'):yield 'heading',clean_text(line.lstrip('#').strip())
        else:yield 'paragraph',clean_text(line)

def export_resources(owner,rows,fmt):
    allowed=('md','txt','json','docx','pdf','zip','original','csv')
    if fmt not in allowed:raise ValueError('支持 md/txt/json/docx/pdf/zip/original/csv')
    eid=uid();folder=ROOT/'exports';folder.mkdir(exist_ok=True)
    content='\n\n---\n\n'.join('# '+r['title']+'\n\n'+r['content'] for r in rows)
    path=folder/(eid+'.'+('zip' if fmt=='original' else fmt))
    if fmt in ('md','txt'):path.write_text(content,'utf-8')
    elif fmt=='json':path.write_text(dumps([{k:r[k] for k in ('id','title','content','source','created')} for r in rows]),'utf-8')
    elif fmt=='csv':
        import csv
        with path.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.writer(f);w.writerow(['资料ID','标题','类型','创建时间','完整内容'])
            for r in rows:w.writerow([r['id'],r['title'],r['scene'],r['created'],r['content']])
    elif fmt=='docx':
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt,RGBColor
        d=Document();style=d.styles['Normal'];style.font.name='Microsoft YaHei';style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        style.font.size=Pt(10.5)
        for name in ('Title','Heading 1','Heading 2'):
            d.styles[name].font.color.rgb=RGBColor(0,0,0);d.styles[name].font.name='Microsoft YaHei';d.styles[name].element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        for r in rows:
            d.add_heading(r['title'],0)
            for kind,value in blocks(r['content']):
                if kind=='heading':d.add_heading(value,2)
                elif kind=='table':
                    width=max(len(x) for x in value);table=d.add_table(rows=0,cols=width);table.style='Table Grid'
                    for row in value:
                        cells=table.add_row().cells
                        for j,cell in enumerate(row):cells[j].text=cell
                    for cell in table.rows[0].cells:
                        for para in cell.paragraphs:
                            for run in para.runs:run.bold=True
                else:d.add_paragraph(value)
        d.save(path)
    elif fmt=='pdf':
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib import colors
        pdfmetrics.registerFont(TTFont('LocalChinese',str(ROOT/'fonts'/'simhei.ttf')))
        body=ParagraphStyle('body',fontName='LocalChinese',fontSize=10.5,leading=17,wordWrap='CJK',spaceAfter=6)
        heading=ParagraphStyle('heading',parent=body,fontSize=13,leading=19,spaceBefore=10,spaceAfter=7,keepWithNext=True)
        title=ParagraphStyle('title',parent=heading,fontSize=18,leading=25)
        story=[]
        for r in rows:
            story.append(Paragraph(html.escape(r['title']),title))
            for kind,value in blocks(r['content']):
                if kind=='table':
                    width=max(len(x) for x in value)
                    cells=[[Paragraph(html.escape(x),body) for x in row]+['']*(width-len(row)) for row in value]
                    table=Table(cells,colWidths=[(595.28-90)/width]*width,repeatRows=1,hAlign='LEFT')
                    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.4,colors.HexColor('#cccccc')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eeeeee')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6)]))
                    story.extend([table,Spacer(1,8)])
                else:story.append(Paragraph(html.escape(value),heading if kind=='heading' else body))
        SimpleDocTemplate(str(path),leftMargin=45,rightMargin=45,topMargin=40,bottomMargin=40).build(story)
    else:
        count=0;manifest=[]
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for r in rows:
                prefix=r['id'][:8]+'/'
                if fmt!='original':z.writestr(prefix+'正文.md',r['content']);z.writestr(prefix+'来源.txt',r['source'])
                for f in json.loads(r['files']):
                    if f.get('status')=='saved':z.write(f['path'],prefix+Path(f['path']).name);count+=1
                manifest.append({'resource_id':r['id'],'title':r['title'],'files':[{k:v for k,v in f.items() if k!='path'} for f in json.loads(r['files'])]})
            z.writestr('manifest.json',dumps(manifest))
        if fmt=='original' and not count:raise ValueError('选中资料没有已保存的原始附件')
    expires=int(time.time())+3600
    with conn() as c:c.execute('INSERT INTO exports VALUES(?,?,?,?)',(eid,owner,str(path),expires))
    return {'download_url':f'{BASE}/download/{eid}?expires={expires}&sig={sign(eid,expires)}','expires_in_seconds':3600,'format':fmt,'export_id':eid}

def calculate(expr):
    if len(expr)>300:raise ValueError('expression too long')
    ops={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv,ast.Pow:operator.pow,ast.Mod:operator.mod}
    def ev(n,depth=0):
        if depth>15:raise ValueError('too complex')
        if isinstance(n,ast.Constant) and type(n.value) in (int,float):v=n.value
        elif isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.UAdd,ast.USub)):v=ev(n.operand,depth+1)*(1 if isinstance(n.op,ast.UAdd) else -1)
        elif isinstance(n,ast.BinOp) and type(n.op) in ops:
            a,b=ev(n.left,depth+1),ev(n.right,depth+1)
            if isinstance(n.op,ast.Pow) and abs(b)>20:raise ValueError('power too large')
            v=ops[type(n.op)](a,b)
        else:raise ValueError('unsupported expression')
        if type(v) not in (int,float) or not math.isfinite(v) or abs(v)>1e100:raise ValueError('result is not a bounded real number')
        return v
    return ev(ast.parse(expr,mode='eval').body)

@app.post('/api/context')
def context():
    p=payload();state=load_state(p['owner'],p['session_id'])
    return jsonify({'state':state,'capabilities':{'map':False,'device':False,'voice':False,'archive':True,'knowledge_sync':True,'export':['md','txt','json','docx','pdf','zip','original','csv']}})

@app.post('/api/exam-check')
def exam_check():
    p=payload();draft=p.get('draft',{});items=draft.get('calculations',[]) if isinstance(draft,dict) else []
    if not isinstance(items,list):items=[]
    checks=[]
    for item in items[:12]:
        if not isinstance(item,dict):continue
        expr=item.get('expression','');row={'label':str(item.get('label',''))[:160],'expression':str(expr)[:300]}
        try:
            if not isinstance(expr,str):raise ValueError('expression must be a string')
            row.update(result=calculate(expr),status='ok')
        except Exception:row.update(status='error',error='算式不可计算，请核对表达式、分母和数值范围')
        checks.append(row)
    return jsonify({'checks':checks,'grid_checks':check_grids(draft.get('binary_grid') if isinstance(draft,dict) else None),'scope':'只核验算术和指定黑白格运算，不证明取数、题意或视觉识别正确。','reference_lookup':False})

@app.post('/api/prepare')
def prepare():
    p=payload();owner=p['owner'];sid=p['session_id'];rid=p['request_id'];route=p.get('route',{});q=p.get('query','')
    with LOCK:
        cached=receipt(owner,rid,'prepare')
        if cached:return jsonify(cached)
        state=load_state(owner,sid);scene=route.get('scene','general');action=route.get('action','answer')
        if scene not in ('general','navigation','homework','meeting','resource'):scene='general'
        result={'scene':scene,'direct_text':'','context':{},'archive':False,'status':'completed','state':state}
        if p.get('input_error'):result['direct_text']=p['input_error'];result['status']='need_input'
        elif route.get('clarification'):result['direct_text']=str(route['clarification']);result['status']='need_input'
        elif action=='stop_playback':result['direct_text']='已收到停止播报事件。本机版尚未连接眼镜播放控制，请在设备上停止播放。';result['status']='unavailable'
        elif action in ('confirm','cancel'):result['direct_text']='当前没有可确认或取消的待办动作。';result['status']='need_input'
        elif action=='capture' and not p.get('files'):result['direct_text']='本机版尚未连接眼镜摄像头，请先上传图片再继续。';result['status']='need_input'
        elif scene=='navigation':
            dest=route.get('destination') or state.get('destination','');mode=route.get('travel_mode') or state.get('travel_mode','步行')
            state.update(destination=dest,travel_mode=mode)
            result['direct_text']=('已记住目的地：'+dest+'，出行方式：'+mode+'。\n\n' if dest else '')+'本机版尚未接入定位和地图服务，暂时不能计算真实路线、距离或耗时。'+('请先告诉我目的地。' if not dest else '接入地图服务后可继续规划。')
            result['status']='unavailable'
        elif scene=='resource':
            try:
                ids=route.get('resource_ids') or (state.get('recent_resources',[])[:1] if action=='export' and any(w in q for w in ('刚才','这份','上面','最新','导出')) else None)
                rows=resource_rows(owner,route.get('search_query',''),ids)
                if not ids:
                    if '纪要' in q:rows=[r for r in rows if r['scene']=='meeting']
                    elif '题目' in q or '题解' in q:rows=[r for r in rows if r['scene']=='homework']
                    if any(w in q for w in ('最新','最近一份','上一份')):rows=rows[:1]
                if action=='export':
                    if not rows:raise ValueError('没有找到可导出的资料，请先生成纪要、题解或保存图片。')
                    if not ids and len(rows)>1:raise ValueError('找到多份资料，请指定资料 ID，或说导出最新一份。')
                    exported=export_resources(owner,rows,route.get('export_format') or 'pdf')
                    result['direct_text']=f"已导出 {len(rows)} 份资料（{exported['format'].upper()}）。\n\n[下载文件]({exported['download_url']})\n\n链接有效期 1 小时；本机访问。"
                    result['context']=exported
                else:
                    result['direct_text']='没有找到相关资料。' if not rows else '\n\n'.join(f"**{r['title']}**\n资料 ID：`{r['id']}`\n创建：{r['created']}\n\n{r['content'][:1600]}" for r in rows[:5])
                    state['recent_resources']=[r['id'] for r in rows[:10]]
            except ValueError as e:result['direct_text']=str(e);result['status']='need_input'
        elif scene=='meeting':
            m=state.setdefault('meeting',{'status':'idle','transcript':''})
            if action=='start':
                if m.get('status')=='collecting':result['direct_text']='当前已有会议文本记录会话，可继续提交转写内容。'
                else:m.update(id=uid(),status='collecting',transcript='');result['direct_text']='已创建会议文本记录会话。当前未连接麦克风，不会自动录音；请粘贴会议转写内容，我会据此生成纪要。'
            else:
                transcript=route.get('transcript') or p.get('transcript','')
                if transcript:
                    if len(m.get('transcript',''))+len(transcript)>48000:result['direct_text']='本机版单场会议文本上限为 48000 字符，请先导出当前会议后新建会议。';result['status']='need_input'
                    else:
                        m.setdefault('id',uid());m['transcript']=(m.get('transcript','')+'\n'+transcript).strip();m['status']='collecting'
                if action=='stop':m['status']='ended'
                if not result['direct_text']:
                    if not m.get('transcript') and not p.get('visual',''):result['direct_text']='还没有会议内容，请粘贴转写文本或上传白板图片。';result['status']='need_input'
                    elif action=='append':result['direct_text']='已保存这段会议文本。说“生成会议纪要”即可汇总。'
                    else:result['context']={'transcript':m.get('transcript',''),'meeting_id':m.get('id'),'status':m.get('status')};result['archive']=True
        elif scene=='homework':
            result['context']={'previous_question':state.get('last_question',''),'previous_answer':state.get('last_answer','')}
            expr=route.get('expression','')
            if expr:
                try:result['context']['calculation']={'expression':expr,'result':calculate(expr)}
                except Exception as e:result['context']['calculation']={'error':str(e)}
            result['archive']=True
        if action=='switch_scene':result['direct_text']='已切换到'+{'homework':'拍题','navigation':'导航','meeting':'会议','general':'普通问答','resource':'资料查询'}[scene]+'模式。';result['archive']=False
        if scene!='resource':state['active_scene']=scene
        save_state(owner,sid,state)
        return jsonify(remember(owner,rid,'prepare',result))

@app.post('/api/finalize')
def finalize():
    p=payload();owner=p['owner'];sid=p['session_id'];rid=p['request_id']
    with LOCK:
        cached=receipt(owner,rid,'finalize')
        if cached:return jsonify(cached)
        prepared=p['prepared'];answer=p['answer'];state=load_state(owner,sid);resource_id=''
        if (prepared.get('archive') or p.get('files')) and answer.strip():
            scene=prepared['scene'];source=p.get('visual') or p.get('query','')
            if scene=='meeting':source=state.get('meeting',{}).get('transcript','')+'\n'+p.get('visual','')
            title={'meeting':'会议纪要','homework':'题目与解答'}.get(scene,'拍摄图片')+' '+datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
            resource_id,files=archive(owner,scene,title,answer,source,p.get('files',[]))
            state['recent_resources']=([resource_id]+state.get('recent_resources',[]))[:20]
            answer+='\n\n---\n已保存到本机资料库，知识库索引已排队。资料 ID：`'+resource_id+'`。可说“导出成 Word/PDF”。'
            if any(f.get('status')!='saved' for f in files):answer+='\n部分原图保存失败；文字结果已保存。'
            if scene=='homework':
                if p.get('files') or not state.get('last_question') or not re.search('为什么|解释|步骤|继续|再讲|换种',p.get('query','')):state['last_question']=source[:12000]
                state['last_answer']=p['answer'][:12000]
        save_state(owner,sid,state)
        result={'answer':answer,'state_json':dumps({k:v for k,v in state.items() if k!='meeting'}),'active_scene':state.get('active_scene','general'),'resource_id':resource_id,'status':prepared.get('status','completed')}
        return jsonify(remember(owner,rid,'finalize',result))

@app.get('/health')
def health():return jsonify({'ok':True,'version':'1.0','mode':'local','dataset_id':SET['dataset_id']})
@app.get('/download/<eid>')
def download(eid):
    try:expiry=int(request.args.get('expires','0'))
    except ValueError:abort(403)
    if expiry<time.time() or not hmac.compare_digest(request.args.get('sig',''),sign(eid,expiry)):abort(403)
    with conn() as c:r=c.execute('SELECT * FROM exports WHERE id=? AND expires=?',(eid,expiry)).fetchone()
    if not r:abort(404)
    return send_file(r['path'],as_attachment=True,download_name='眼镜助手资料'+Path(r['path']).suffix)

if __name__=='__main__':app.run(host='0.0.0.0',port=8792,threaded=True)
